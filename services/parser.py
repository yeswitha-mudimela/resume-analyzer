"""
Multi-format resume document parser.
Extracts raw text, structural metadata, and inspects document health across
PDF, DOCX, and plain text formats.
"""

import io
import re
from typing import BinaryIO, Dict, Any, Tuple

# Optional and required document libraries
try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import docx
except ImportError:
    docx = None


class ParseError(Exception):
    """Custom exception raised when resume text extraction fails."""
    pass


def extract_text_from_pdf(file_stream: BinaryIO) -> Tuple[str, int]:
    """
    Extracts text from a PDF stream using pdfplumber.
    Returns a tuple of (extracted_text, page_count).
    """
    if pdfplumber is None:
        raise ParseError("pdfplumber is not installed on the server.")

    try:
        file_stream.seek(0)
        text_parts = []
        with pdfplumber.open(file_stream) as pdf:
            page_count = len(pdf.pages)
            if page_count == 0:
                raise ParseError("The uploaded PDF has 0 pages.")

            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)

        full_text = "\n".join(text_parts).strip()
        return full_text, page_count
    except Exception as e:
        raise ParseError(f"Failed to parse PDF document: {str(e)}")


def extract_text_from_docx(file_stream: BinaryIO) -> Tuple[str, int]:
    """
    Extracts text from a DOCX stream using python-docx.
    Estimates page count based on typical word density (~350 words/page).
    """
    if docx is None:
        raise ParseError("python-docx is not installed on the server.")

    try:
        file_stream.seek(0)
        doc = docx.Document(file_stream)
        text_parts = []

        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_parts.append(" | ".join(row_text))

        full_text = "\n".join(text_parts).strip()
        word_count = len(full_text.split())
        estimated_pages = max(1, round(word_count / 350))
        return full_text, estimated_pages
    except Exception as e:
        raise ParseError(f"Failed to parse DOCX document: {str(e)}")


def extract_text_from_txt(file_stream: BinaryIO) -> Tuple[str, int]:
    """Extracts text from a plain text file stream."""
    try:
        file_stream.seek(0)
        raw = file_stream.read()
        try:
            full_text = raw.decode("utf-8").strip()
        except UnicodeDecodeError:
            full_text = raw.decode("latin-1", errors="replace").strip()

        word_count = len(full_text.split())
        estimated_pages = max(1, round(word_count / 350))
        return full_text, estimated_pages
    except Exception as e:
        raise ParseError(f"Failed to parse text file: {str(e)}")


def parse_resume_document(file_stream: BinaryIO, filename: str) -> Dict[str, Any]:
    """
    Main entry point for parsing any supported resume file.
    Inspects extension, extracts text, computes metadata, and detects scanned/image files.
    """
    if not filename:
        raise ParseError("No filename provided.")

    lower_name = filename.lower()

    if lower_name.endswith(".pdf"):
        text, pages = extract_text_from_pdf(file_stream)
        file_type = "pdf"
    elif lower_name.endswith(".docx"):
        text, pages = extract_text_from_docx(file_stream)
        file_type = "docx"
    elif lower_name.endswith(".txt"):
        text, pages = extract_text_from_txt(file_stream)
        file_type = "txt"
    else:
        raise ParseError("Unsupported file type. Please upload a .pdf, .docx, or .txt file.")

    # Calculate metrics
    words = text.split()
    word_count = len(words)
    char_count = len(text)
    lines = [line for line in text.split("\n") if line.strip()]
    line_count = len(lines)

    # Scanned document detection: only applies to PDFs where pages exist but extracted text is negligible
    is_scanned = bool(file_type == "pdf" and pages > 0 and word_count < 25)

    # Reading time calculation (average 200 words per minute)
    reading_time_minutes = max(1, round(word_count / 200, 1))

    return {
        "text": text,
        "filename": filename,
        "file_type": file_type,
        "page_count": pages,
        "word_count": word_count,
        "char_count": char_count,
        "line_count": line_count,
        "is_scanned": is_scanned,
        "reading_time_minutes": reading_time_minutes
    }
