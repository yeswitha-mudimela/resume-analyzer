import io
import pytest
from services.parser import (
    parse_resume_document,
    extract_text_from_txt,
    extract_text_from_docx,
    ParseError
)

def test_extract_text_from_txt():
    sample_text = "John Doe\nSoftware Engineer\nPython, SQL, Docker\nExperience at Tech Corp"
    stream = io.BytesIO(sample_text.encode("utf-8"))
    text, pages = extract_text_from_txt(stream)
    assert "John Doe" in text
    assert "Software Engineer" in text
    assert pages >= 1

def test_parse_resume_document_txt():
    content = "Alice Smith\nFrontend Developer\nSkills: React, JavaScript, HTML, CSS\nEmail: alice@example.com"
    stream = io.BytesIO(content.encode("utf-8"))
    result = parse_resume_document(stream, "alice_resume.txt")
    
    assert result["file_type"] == "txt"
    assert result["word_count"] > 5
    assert result["char_count"] > 20
    assert result["is_scanned"] is False
    assert "alice@example.com" in result["text"]

def test_parse_unsupported_file_extension():
    stream = io.BytesIO(b"fake image data")
    with pytest.raises(ParseError) as exc_info:
        parse_resume_document(stream, "resume.png")
    assert "Unsupported file type" in str(exc_info.value)

def test_parse_resume_document_docx():
    import docx
    doc = docx.Document()
    doc.add_heading("Bob Johnson", 0)
    doc.add_paragraph("Full Stack Developer with 5 years experience.")
    doc.add_paragraph("Skills: Python, TypeScript, React, PostgreSQL")
    
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    
    result = parse_resume_document(stream, "bob_resume.docx")
    assert result["file_type"] == "docx"
    assert "Bob Johnson" in result["text"]
    assert "PostgreSQL" in result["text"]
    assert result["word_count"] > 10
